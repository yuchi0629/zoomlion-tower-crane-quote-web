"""Build the static GitHub Pages data files from the desktop database.

The browser never imports Excel files. This tool applies the desktop import
rules, normalizes the database into public JSON, and embeds the already-filtered
two-sheet configuration workbook used by the browser export.
"""

from __future__ import annotations

import argparse
import base64
import copy
import json
import sys
import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


HERE = Path(__file__).resolve()
WEB_ROOT = HERE.parents[1]
DEFAULT_DESKTOP_ROOT = HERE.parents[3]


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def camel_form(item):
    return {
        "installCode": clean(item.get("install_code")),
        "installForm": clean(item.get("install_form")),
        "height": clean(item.get("height")),
        "jibLength": clean(item.get("jib_length")),
        "ropeCapacity": clean(item.get("rope_capacity")),
        "wireRopeSpec": clean(item.get("wire_rope_spec")),
        "fall": clean(item.get("fall")),
        "hoistingHeight": clean(item.get("hoisting_height")),
    }


def camel_price_option(item):
    return {
        "categoryCode": clean(item.get("category_code")),
        "category": clean(item.get("category")),
        "partCode": clean(item.get("part_code")),
        "partName": clean(item.get("part_name")),
        "materialCode": clean(item.get("material_code")),
        "materialDescription": clean(item.get("material_description")),
        "changeTypeCode": clean(item.get("change_type_code")),
        "addPrice": clean(item.get("add_price")),
        "deductPrice": clean(item.get("deduct_price")),
    }


def camel_product(item):
    tower_type = item.get("tower_type") or {}
    return {
        "model": clean(item.get("model")),
        "productCode": clean(item.get("product_code")),
        "configDescription": clean(item.get("config_description")),
        "towerSize": clean(item.get("tower_size")),
        "defaultJibLength": clean(item.get("default_jib_length")),
        "ratedTonnage": clean(item.get("rated_tonnage")),
        "mastType": clean(item.get("mast_type")),
        "towerType": {
            "zh": clean(tower_type.get("zh")),
            "en": clean(tower_type.get("en")),
        },
        "maxLoad": clean(item.get("max_load")),
    }


def overlay_price_workbook(db, path, app):
    if not path or not path.exists():
        return
    ws = load_workbook(path, data_only=True, read_only=True).active
    replacement_options = {}
    seen = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if clean(row[3] if len(row) > 3 else "") != "B1":
            continue
        model = clean(row[4] if len(row) > 4 else "")
        description = clean(row[5] if len(row) > 5 else "")
        product_code = clean(row[6] if len(row) > 6 else "")
        if not model or not product_code:
            continue
        parsed = app.parse_config_description(description)
        db["products"][model] = {
            "model": model,
            "model_key": app.normalize_key(model),
            "product_code": product_code,
            "config_description": description,
            "tower_size": parsed["tower_size"],
            "default_jib_length": parsed["jib_length"],
            "rated_tonnage": parsed["tonnage"],
            "mast_type": parsed["mast_type"],
            "tower_type": app.detect_tower_type(model),
            "max_load": app.detect_max_load(model),
        }
        option = {
            "category_code": clean(row[8] if len(row) > 8 else ""),
            "category": clean(row[9] if len(row) > 9 else ""),
            "part_code": clean(row[10] if len(row) > 10 else ""),
            "part_name": clean(row[11] if len(row) > 11 else ""),
            "material_code": clean(row[12] if len(row) > 12 else ""),
            "material_description": clean(row[13] if len(row) > 13 else ""),
            "change_type_code": clean(row[14] if len(row) > 14 else ""),
            "add_price": app.format_number(row[15] if len(row) > 15 else ""),
            "deduct_price": app.format_number(row[16] if len(row) > 16 else ""),
        }
        if not (option["part_name"] or option["part_code"]):
            continue
        dedupe = "|".join(
            [
                option["category_code"],
                option["part_code"],
                option["part_name"],
                option["change_type_code"],
            ]
        )
        if dedupe in seen.setdefault(product_code, set()):
            continue
        seen[product_code].add(dedupe)
        replacement_options.setdefault(product_code, []).append(option)
    for product_code, options in replacement_options.items():
        db["options_by_code"][product_code] = options


def overlay_form_workbook(db, path, app):
    if not path or not path.exists():
        return
    ws = load_workbook(path, data_only=True, read_only=True).active
    for row in ws.iter_rows(min_row=2, values_only=True):
        product_code = clean(row[0] if len(row) > 0 else "")
        install_form = clean(row[2] if len(row) > 2 else "")
        if not product_code or not install_form:
            continue
        form = {
            "install_code": clean(row[1] if len(row) > 1 else ""),
            "install_form": install_form,
            "height": clean(row[9] if len(row) > 9 else ""),
            "jib_length": clean(row[10] if len(row) > 10 else ""),
            "rope_capacity": clean(row[11] if len(row) > 11 else ""),
            "wire_rope_spec": clean(row[12] if len(row) > 12 else ""),
            "fall": clean(row[13] if len(row) > 13 else ""),
            "hoisting_height": clean(row[14] if len(row) > 14 else ""),
        }
        forms = db["forms_by_code"].setdefault(product_code, [])
        form_key = app.normalize_form(install_form)
        forms[:] = [
            item
            for item in forms
            if app.normalize_form(item.get("install_form")) != form_key
        ]
        forms.append(form)


def row_payload(row):
    return {
        "composition": clean(row.get("composition")),
        "component": clean(row.get("component")),
        "name": clean(row.get("name")),
        "code": clean(row.get("code")),
        "modelCode": clean(row.get("model_code")) or "/",
        "mark": clean(row.get("mark")),
        "price": clean(row.get("price")),
        "itemDisplay": clean(row.get("item_display")),
    }


def option_payload(item):
    result = row_payload(item)
    result["children"] = [row_payload(child) for child in item.get("children") or []]
    return result


def main_component_payload(item):
    return {
        "component": clean(item.get("component_zh")),
        "componentEn": clean(item.get("component_en")),
        "name": clean(item.get("part_name")),
        "code": clean(item.get("code")) or "/",
        "quantity": clean(item.get("quantity")) or "1",
    }


def workbook_metadata(app, db, model):
    workbook_bytes = app.workbook_bytes_for_model(db, model)
    if not workbook_bytes:
        return {}, ""
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    basic_ws = app.worksheet_for_list_type(workbook, "basic")
    option_ws = app.worksheet_for_list_type(workbook, "option")
    basic_start, _ = app.list_section_bounds(basic_ws, "basic")
    option_start, _ = app.list_section_bounds(option_ws, "option")
    metadata = {
        "basicTitle": clean(basic_ws.cell(basic_start, 1).value),
        "basicVersion": clean(basic_ws.cell(basic_start + 1, 1).value),
        "optionTitle": clean(option_ws.cell(option_start, 1).value),
        "optionVersion": clean(option_ws.cell(option_start + 1, 1).value),
    }
    workbooks = {}
    with tempfile.TemporaryDirectory() as temp_dir:
        for language in ("zh", "en", "fr", "de"):
            output = app.make_combined_config_option_excel(
                db,
                model,
                Path(temp_dir) / f"combined-{language}.xlsx",
                language,
            )
            exported = load_workbook(output)
            brand_fill = PatternFill(
                fill_type="solid",
                fgColor="AADB1E",
                start_color="AADB1E",
                end_color="AADB1E",
            )
            for worksheet in exported.worksheets:
                for row_index in (1, 3, 4):
                    for cell in worksheet[row_index]:
                        cell.font = Font(
                            name="Arial",
                            size=10,
                            bold=True,
                            color=cell.font.color,
                        )
                for row_index in (3, 4):
                    for cell in worksheet[row_index]:
                        cell.fill = brand_fill
            exported.save(output)
            workbooks[language] = base64.b64encode(Path(output).read_bytes()).decode(
                "ascii"
            )
    return metadata, workbooks


def build_product(app, db, model, published):
    product = db["products"][model]
    product_code = clean(product.get("product_code"))
    result = camel_product(product)
    result["published"] = published
    result["priceOptions"] = [
        camel_price_option(item)
        for item in db.get("options_by_code", {}).get(product_code, [])
    ]
    result["forms"] = []
    metadata = {}
    combined_workbooks = {}
    if published:
        metadata, combined_workbooks = workbook_metadata(app, db, model)
    for raw_form in db.get("forms_by_code", {}).get(product_code, []):
        form = camel_form(raw_form)
        install_form = form["installForm"]
        form.update(
            {
                "machinePrice": "0",
                "basicRows": [],
                "optionRows": [],
                "mainComponents": [],
            }
        )
        if published:
            basic_rows = app.basic_config_items(db, model, install_form)
            option_rows = app.option_config_items(db, model, install_form)
            components = app.select_components(db, model, install_form)
            form["basicRows"] = [row_payload(item) for item in basic_rows]
            form["optionRows"] = [option_payload(item) for item in option_rows]
            form["mainComponents"] = [
                main_component_payload(item) for item in components
            ]
            form["machinePrice"] = next(
                (
                    clean(item.get("price"))
                    for item in basic_rows
                    if app.parse_price_value(item.get("price"))
                ),
                "0",
            )
        result["forms"].append(form)
    result["listMetadata"] = metadata
    if combined_workbooks:
        result["combinedWorkbooks"] = combined_workbooks
    return result


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--desktop-root", type=Path, default=DEFAULT_DESKTOP_ROOT)
    parser.add_argument("--price-file", type=Path)
    parser.add_argument("--form-file", type=Path)
    parser.add_argument("--output-dir", type=Path, default=WEB_ROOT / "public" / "data")
    args = parser.parse_args()

    desktop_root = args.desktop_root.resolve()
    sys.path.insert(0, str(desktop_root))
    import app

    database_path = desktop_root / "data" / "quotation_database.json"
    db = json.loads(database_path.read_text(encoding="utf-8"))
    price_file = args.price_file or (
        desktop_root
        / "参考文件"
        / "PriceInDecreasePrice_dateOut_2026062107481603503417_1_1.xlsx"
    )
    form_file = args.form_file or (
        desktop_root
        / "参考文件"
        / "PrdTowerCraneCfg_dataOut_2026062107463298884388_1_1.xlsx"
    )
    overlay_price_workbook(db, price_file, app)
    overlay_form_workbook(db, form_file, app)

    published_keys = {
        app.normalize_key(item)
        for item in db.get("published_config_model_keys", [])
        if clean(item)
    }
    visible_models = [
        model
        for model in db.get("products", {})
        if app.is_visible_product_model(model)
    ]
    visible_models.sort(
        key=lambda model: (
            0 if app.normalize_key(model) in published_keys else 1,
            model,
        )
    )
    products = [
        build_product(
            app,
            db,
            model,
            app.normalize_key(model) in published_keys,
        )
        for model in visible_models
    ]

    existing_path = args.output_dir / "app-data.json"
    existing = (
        json.loads(existing_path.read_text(encoding="utf-8"))
        if existing_path.exists()
        else {}
    )
    app_data = {
        "version": "1.3.0",
        "generatedFrom": [
            database_path.name,
            price_file.name,
            form_file.name,
        ],
        "brandColor": "#AADB1E",
        "dataSummary": {
            "visibleProducts": len(products),
            "allProducts": len(db.get("products", {})),
            "installationForms": sum(
                len(items) for items in db.get("forms_by_code", {}).values()
            ),
            "priceOptions": sum(
                len(items) for items in db.get("options_by_code", {}).values()
            ),
            "publishedConfigurations": [
                item["model"] for item in products if item["published"]
            ],
        },
        "products": products,
        "translations": existing.get("translations", {}),
        "ui": existing.get("ui", {}),
    }
    write_json(args.output_dir / "app-data.json", app_data)

    product_catalog = [camel_product(item) for item in db.get("products", {}).values()]
    form_catalog = [
        {"productCode": product_code, **camel_form(item)}
        for product_code, items in db.get("forms_by_code", {}).items()
        for item in items
    ]
    option_catalog = [
        {"productCode": product_code, **camel_price_option(item)}
        for product_code, items in db.get("options_by_code", {}).items()
        for item in items
    ]
    write_json(args.output_dir / "products.json", product_catalog)
    write_json(args.output_dir / "install-forms.json", form_catalog)
    write_json(args.output_dir / "price-options.json", option_catalog)
    print(json.dumps(app_data["dataSummary"], ensure_ascii=False))


if __name__ == "__main__":
    main()
